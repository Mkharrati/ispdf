import os
import img2pdf
import pikepdf
import aspose.words as aw
import docx2pdf
import docx
import pptx.presentation
from pypdf import PdfWriter
from pypdf import PdfReader
import file_utils
import pdf2image
import pptxtopdf
import pptx
import requests
import json
import API.OCR as ocr

OCR_URL = "https://www.eboo.ir/api/ocr/getway"

def convert_images_to_pdf(user_folder, output_pdf):
    """
    Convert all images in a folder to a PDF.
    """
    A4_size = (img2pdf.mm_to_pt(210),img2pdf.mm_to_pt(297))
    layout_fun = img2pdf.get_layout_fun(A4_size)
    image_files = file_utils.list_files_by_time(user_folder)
    with open(output_pdf, "wb") as f:
        f.write(img2pdf.convert(image_files, layout_fun=layout_fun))
    return output_pdf

def convert_pdf_to_docx_Aspose_Words(pdf_path, output_docx):
    """
    Convert a PDF file to a DOCX file using Aspose.Words.
    """
    document = aw.Document(pdf_path)
    document.save(output_docx)
    os.remove(pdf_path)
    return output_docx

def convert_docx_to_pdf(docx_path, output_pdf):
    """
    Convert a DOCX file to PDF using docx2pdf.
    """
    docx2pdf.convert(docx_path, output_pdf)
    return output_pdf

def merge_pdfs(pdf_files, output_pdf):
    """
    Merge multiple PDF files into a single PDF.
    """
    merger = PdfWriter()
    for pdf in pdf_files:
        merger.append(pdf)
    merger.write(output_pdf)
    merger.close()
    return output_pdf

def unlock_pdf(pdf_path):
    """
    Remove password restrictions from a PDF using pikepdf.
    """
    try:
        pdf = pikepdf.open(pdf_path)
        unlocked_path = f"{pdf_path}_Unlocked.pdf"
        pdf.save(unlocked_path)
        pdf.close()
        return unlocked_path
    except Exception as e:
        print(f"Error unlocking PDF: {e}")
        return None

def is_pdf_file(pdf_path):
    """Check the file that is a valid PDF."""
    if not pdf_path.endswith(".pdf"):
        return False
    try:
        pikepdf.open(pdf_path)
        return True
    except Exception:
        return False

def is_docx_file(docx_path):
    """Check if a file is a valid DOCX file."""
    try:
        docx.Document(docx_path)
        return True
    except Exception:
        return False
    
def is_pptx_file(pptx_path):
    try:
        pptx.Presentation(pptx_path)
        return True
    except:
        return False
    
def convert_pdf_to_image(user_folder, pdf_path):
    """
    Convert pdf pages to image and save in user_folder path.
    """
    pages = pdf2image.convert_from_path(pdf_path)
    for page in pages:
        random_name = file_utils.random_name()
        image_path = f"{user_folder}/{random_name}.jpg"
        page.save(image_path, "JPEG")
    os.remove(pdf_path)
    return file_utils.list_files_by_time(user_folder)

def convert_pptx_to_pdf(pptx_path, pdf_folder_path):
    """
    Convert Powerpoint pages to pdf file
    """
    pptxtopdf.convert(pptx_path, pdf_folder_path)

def number_of_pdf_pages(pdf_path):
    """
    return the number of the pdf pages via pdf_path.
    """
    reader = PdfReader(pdf_path)
    num_page = len(reader.pages)
    return num_page

def convert_pdf_to_docx(pdf_path, output_docx, OCR_TOKEN):
    """
    convert pdf to docx via api. return docx path
    """
    json_data = ocr.send_file(pdf_path, OCR_TOKEN)
    if "False" in json_data:
        print(json_data[1]) # for debug
        return "False"
    file_token = json_data["FileToken"]
    docx_url = ocr.pdf_to_docx(file_token, OCR_TOKEN)
    if "False" in docx_url:
        print(docx_url[1]) # for debug
        return "False"
    file_content = file_utils.download_link(docx_url)
    print("docx downloaded from server successfully")
    docx_path = file_utils.save_file(file_content, output_docx)
    return docx_path

def image_to_text(image_path, OCR_TOKEN):
    """
    extract text from image with api. return extracted text
    """
    json_data = ocr.send_file(image_path, OCR_TOKEN)
    if "False" in json_data:
        print(json_data[1]) # for debug
        return "False"
    file_token = json_data["FileToken"]
    text = ocr.image_to_text(file_token, OCR_TOKEN)
    return text